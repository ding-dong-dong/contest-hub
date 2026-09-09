#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Excel 竞赛数据 -> JSON（映射到竞赛数据模型）。

数据源：文博卉第一天交付.xlsx 的「全量明细表」sheet（表头在第 1 行）。
输出：contests_import.json，字段与 app/schemas.py 的 ContestCreate 对齐，
可直接用于 POST /contests 批量导入（其中 id 为 imp_序号，可重复执行不冲突）。

用法：
    python scripts/import_excel.py
    python scripts/import_excel.py --input 其他文件.xlsx --output out.json
"""
import argparse
import calendar
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("缺少依赖：请先执行 pip install openpyxl")

# ---- 源 Excel 列名（与「全量明细表」表头严格一致）----
COL_SEQ = "序号"
COL_NAME = "竞赛名称"
COL_EDITION = "本届届次"
COL_LEVEL = "比赛等级"
COL_MAJOR_SRC = "适配专业"
COL_GOAL = "适配理想/目标"
COL_QUAL = "详细参赛资格"
COL_PERIOD = "比赛周期(本届)"
COL_DEADLINE = "报名截止(本届)"
COL_ENTRY_STATUS = "官方报名入口状态"
COL_NOTICE_URL = "官方通知URL"
COL_REG_URL = "报名链接"
COL_MATERIALS = "需准备材料"
COL_FLOW = "提交流程"
COL_SCHOOL_TYPE = "适用学校类别"
COL_VALUE = "竞赛自身价值(客观)"
COL_CERT = "学校认定说明(已核实才录入)"
COL_TAG_FIRST = "首次参赛"
COL_TAG_GOAL = "已有目标"
COL_TAG_TIME = "时间有限"
COL_TAG_CERT = "重视认定"
COL_VERIFIED = "最近核查"
COL_STATUS = "状态"

GRADE_ORDER = ["大一", "大二", "大三", "大四", "研究生"]


def cell(v) -> str:
    """单元格清洗：None->空串；datetime/date->ISO；其余去首尾空白。"""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).replace("\r\n", "\n").replace("\r", "\n").strip()


def extract_deadline(text: str) -> str:
    """从自由文本中提取报名截止日期，输出 ISO 日期；取不到返回空串。

    支持：2026-09-25 / 2026年9月25日 / 2026年5月底前（取当月最后一天）。
    """
    if not text:
        return ""
    m = re.search(r"(\d{4})\s*-\s*(\d{1,2})\s*-\s*(\d{1,2})", text)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月(?:\s*(\d{1,2})\s*日?)?", text)
    if m:
        y, mo = int(m.group(1)), int(m.group(2))
        d = m.group(3)
        if d:
            day = int(d)
        else:
            tail = text[m.end():m.end() + 6]
            day = calendar.monthrange(y, mo)[1] if ("底" in tail or "末" in tail) else 1
        return f"{y:04d}-{mo:02d}-{day:02d}"
    # 无年份的「报名X月X日—Y月Y日」，取报名窗口结束日，年份按当前年
    m = re.search(
        r"报名[^。；;0-9]{0,8}\d{1,2}月\d{1,2}日?\s*[-—~至到]+\s*(\d{1,2})月(\d{1,2})日",
        text,
    )
    if m:
        y = date.today().year
        return f"{y:04d}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    # 无年份但明确带「报名/截止」前缀的月日
    m = re.search(r"(?:报名|截止)[^。；;0-9]{0,10}(\d{1,2})月(\d{1,2})日", text)
    if m:
        y = date.today().year
        return f"{y:04d}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    return ""


def extract_organizer(value_text: str, qual_text: str) -> str:
    """从「竞赛自身价值」或「详细参赛资格」文本中提取主办单位。

    覆盖句式：
      由XXX联合主办 / XXX共同主办,YYY承办 / XX指导,YYY主办 / XXX主办。
    取「主办」首次出现处之前的最后一个分句，去掉「由/共同/联合/等」等修饰。
    """
    for txt in (value_text, qual_text):
        if not txt:
            continue
        idx = txt.find("主办")
        if idx == -1:
            continue
        seg = re.split(r"[，,；;。！!？?]", txt[:idx])[-1]
        seg = re.sub(r"^由\s*", "", seg).strip()
        seg = re.sub(r"(共同|联合|牵头)$", "", seg).strip()
        seg = re.sub(r"等$", "", seg).strip(" 、,，")
        if seg:
            return seg
    return ""


def extract_grades(qual: str) -> list:
    """从参赛资格文本推断适用年级；推断不出但提及本科/在校则默认本科四年。"""
    if not qual:
        return []
    grades = []
    for g in GRADE_ORDER[:-1]:
        if g in qual:
            grades.append(g)
    if re.search(r"大二及以上|大二以上|大二[（(]含[）)]", qual):
        for g in ["大二", "大三", "大四"]:
            if g not in grades:
                grades.append(g)
    has_postgrad = ("研究生" in qual) or ("硕博" in qual) or ("硕士" in qual and "博士" in qual)
    if not grades and ("本科" in qual or "专科" in qual or "在校" in qual):
        grades = ["大一", "大二", "大三", "大四"]
    if has_postgrad:
        grades.append("研究生")
    return [g for g in GRADE_ORDER if g in grades]


def normalize_major(text: str) -> str:
    if not text:
        return "不限"
    if "任何专业" in text or "不限" in text:
        return "不限"
    # 去掉括号补充说明，如 人文、社科、传媒、理工(综...)
    return re.sub(r"[（(].*?[）)]", "", text).strip() or "不限"


def normalize_school(text: str) -> str:
    if not text:
        return "待确认"
    if text.startswith("全国"):
        return "全国"
    if "指定" in text:
        return "指定院校"
    return "待确认"


def normalize_status(text: str) -> str:
    s = text.strip()
    return {"待复查": "待确认"}.get(s, s or "待确认")


def infer_category(name: str) -> str:
    """按竞赛名称关键词粗略归类，方便筛选；不确定归「综合」。"""
    if any(k in name for k in ("创业", "创新", "互联网+")):
        return "创新创业"
    if any(k in name for k in ("数学", "建模")):
        return "数学建模"
    if "设计" in name:
        return "设计"
    if any(k in name for k in ("电子", "嵌入式", "单片机")):
        return "电子"
    if any(k in name for k in ("人工智能", "AI", "算法", "程序设计", "ACM")):
        return "AI/编程"
    if any(k in name for k in ("学术", "科技作品", "课外学术")):
        return "学术科技"
    if any(k in name for k in ("英语", "演讲", "写作")):
        return "语言文化"
    return "综合"


def build_review_note(r: dict) -> str:
    """未直接映射的字段全部落入 review_note，信息不丢失。"""
    parts = []
    if r[COL_FLOW]:
        parts.append(f"【提交流程】\n{r[COL_FLOW]}")
    extra = [
        ("本届届次", COL_EDITION),
        ("比赛等级", COL_LEVEL),
        ("适配专业(原文)", COL_MAJOR_SRC),
        ("适配理想/目标", COL_GOAL),
        ("比赛周期", COL_PERIOD),
        ("报名入口状态", COL_ENTRY_STATUS),
        ("适用学校类别(原文)", COL_SCHOOL_TYPE),
        ("竞赛自身价值(客观)", COL_VALUE),
        ("学校认定说明", COL_CERT),
        ("详细参赛资格", COL_QUAL),
    ]
    for label, key in extra:
        if r[key]:
            parts.append(f"【{label}】{r[key]}")
    tags = []
    for label, key in (
        ("首次参赛", COL_TAG_FIRST),
        ("已有目标", COL_TAG_GOAL),
        ("时间有限", COL_TAG_TIME),
        ("重视认定", COL_TAG_CERT),
    ):
        if r[key]:
            tags.append(f"{label}：{r[key]}")
    if tags:
        parts.append("【画像标签】" + "；".join(tags))
    return "\n".join(parts)


def convert_row(r: dict) -> dict:
    name = cell(r[COL_NAME])
    notice_url = cell(r[COL_NOTICE_URL])
    deadline_raw = cell(r[COL_DEADLINE])
    seq = cell(r[COL_SEQ])
    seq_num = re.sub(r"\D", "", seq)
    contest = {
        "id": f"imp_{int(seq_num):03d}" if seq_num else None,
        "name": name,
        "category": infer_category(name),
        "organizer": extract_organizer(cell(r[COL_VALUE]), cell(r[COL_QUAL])),
        "eligible_grades": extract_grades(cell(r[COL_QUAL])),
        "major_limit": normalize_major(cell(r[COL_MAJOR_SRC])),
        "school_limit": normalize_school(cell(r[COL_SCHOOL_TYPE])),
        "registration_deadline": extract_deadline(deadline_raw),
        "submission_deadline": "",
        "materials": cell(r[COL_MATERIALS]),
        "skills": "",
        "estimated_time": "",
        "notice_url": notice_url,
        "registration_url": cell(r[COL_REG_URL]),
        "source_type": "官网" if notice_url else "其他",
        "verified_at": extract_deadline(cell(r[COL_VERIFIED])),
        "status": normalize_status(cell(r[COL_STATUS])),
        "review_note": build_review_note(r),
    }
    return contest


def main() -> None:
    base_dir = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description="Excel 竞赛数据转 JSON（映射竞赛数据模型）")
    parser.add_argument(
        "--input",
        default=r"C:\Users\Lenovo\Documents\xwechat_files\wxid_xiyy9d1l0ic222_0726\msg\file\2026-09\文博卉第一天交付.xlsx",
        help="Excel 文件路径",
    )
    parser.add_argument(
        "--output",
        default=str(base_dir / "contests_import.json"),
        help="输出 JSON 路径",
    )
    parser.add_argument("--sheet", default="全量明细表", help="sheet 名称")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.input, data_only=True)
    if args.sheet not in wb.sheetnames:
        sys.exit(f"sheet「{args.sheet}」不存在，可用：{wb.sheetnames}")
    ws = wb[args.sheet]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        sys.exit("sheet 为空")
    header = [cell(c) for c in rows[0]]
    col_idx = {name: i for i, name in enumerate(header)}

    missing = [c for c in (COL_NAME, COL_DEADLINE, COL_STATUS) if c not in col_idx]
    if missing:
        sys.exit(f"表头缺少必要列：{missing}；实际表头：{header}")

    contests = []
    warnings = []
    for row in rows[1:]:
        r = {name: (row[i] if i < len(row) else None) for name, i in col_idx.items()}
        if not cell(r[COL_NAME]):
            continue  # 跳过空行
        c = convert_row(r)
        if not c["registration_deadline"]:
            warnings.append(f"  - {c['name']}：未能从「报名截止」提取日期（原文：{cell(r[COL_DEADLINE])[:40]}）")
        if not c["organizer"]:
            warnings.append(f"  - {c['name']}：未能提取主办单位")
        contests.append(c)

    Path(args.output).write_text(
        json.dumps(contests, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"转换完成：{len(contests)} 条 -> {args.output}")
    if warnings:
        print("提示：")
        print("\n".join(warnings))


if __name__ == "__main__":
    main()
